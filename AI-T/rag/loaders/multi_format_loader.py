"""
rag/loaders/multi_format_loader.py

Loads uploaded documents (.md, .docx, .txt) from a list of file paths
and returns LangChain Document objects ready for the upload pipeline.
"""

import os
import zipfile
import xml.etree.ElementTree as ET
from typing import List
from langchain_core.documents import Document


def _extract_docx_text(file_path: str) -> str:
    """Extract plain text from a .docx file."""
    try:
        import docx
        doc = docx.Document(file_path)
        text = []
        for p in doc.paragraphs:
            if p.text.strip():
                text.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text.append(" | ".join(row_text))
        return "\n\n".join(text)
    except Exception:
        # Fallback: raw XML extraction
        try:
            with zipfile.ZipFile(file_path) as z:
                xml_content = z.read('word/document.xml')
                tree = ET.fromstring(xml_content)
                texts = []
                for elem in tree.iter():
                    if elem.tag.endswith('}t') and elem.text:
                        texts.append(elem.text)
                return " ".join(texts)
        except Exception as e:
            raise RuntimeError(f"Could not extract text from {file_path}: {e}")


class MultiFormatLoader:
    """
    Loads .md, .docx, and .txt files from a provided list of file paths.
    Returns a list of LangChain Document objects tagged with the given account slug and document category.
    """

    def __init__(self, file_paths: List[str], account_slug: str, account_name: str, document_category: str = 'uploaded'):
        """
        Args:
            file_paths:          Absolute paths to the uploaded temporary files.
            account_slug:        Normalized slug (e.g. "acme_corp") used as metadata.account.
            account_name:        Human-readable name (e.g. "Acme Corp") for metadata.
            document_category:   One of: customer | deal_history | seller | uploaded
        """
        self.file_paths = file_paths
        self.account_slug = account_slug
        self.account_name = account_name
        self.document_category = document_category

    def load(self) -> List[Document]:
        documents = []

        for file_path in self.file_paths:
            if not os.path.exists(file_path):
                print(f"[MultiFormatLoader] File not found, skipping: {file_path}")
                continue

            filename = os.path.basename(file_path)
            ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
            content = None

            try:
                if ext == 'md':
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()

                elif ext == 'docx' and not filename.startswith('~$'):
                    content = _extract_docx_text(file_path)

                elif ext == 'txt':
                    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                        content = f.read()

                else:
                    print(f"[MultiFormatLoader] Unsupported file type '{ext}', skipping: {filename}")
                    continue

            except Exception as e:
                print(f"[MultiFormatLoader] Error reading {filename}: {e}")
                continue

            if content and content.strip():
                doc = Document(
                    page_content=content.strip(),
                    metadata={
                        "source": file_path,
                        "filename": filename,
                        "folder": self.document_category,        # customer | deal_history | seller
                        "document_category": self.document_category,
                        "account": self.account_slug,
                        "company": self.account_slug,
                        "account_name": self.account_name,
                        "document_type": ext,
                        "upload_source": "manager_upload"
                    }
                )
                documents.append(doc)
                print(f"[MultiFormatLoader] Loaded: {filename} ({len(content)} chars)")

        print(f"[MultiFormatLoader] Total documents loaded: {len(documents)}")
        return documents
