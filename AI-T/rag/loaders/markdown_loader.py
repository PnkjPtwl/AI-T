import os
import zipfile
import xml.etree.ElementTree as ET
from langchain_core.documents import Document
from typing import List

def extract_docx_text(file_path: str) -> str:
    try:
        import docx
        doc = docx.Document(file_path)
        text = []
        for p in doc.paragraphs:
            if p.text.strip():
                text.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text.append(" | ".join(row_text))
        return "\n\n".join(text)
    except Exception:
        with zipfile.ZipFile(file_path) as z:
            xml_content = z.read('word/document.xml')
            tree = ET.fromstring(xml_content)
            texts = []
            for elem in tree.iter():
                if elem.tag.endswith('}t') and elem.text:
                    texts.append(elem.text)
            return " ".join(texts)

class MarkdownLoader:
    def __init__(self, directory: str):
        self.directory = directory

    def load(self) -> List[Document]:
        documents = []
        for root, dirs, files in os.walk(self.directory):
            dirs[:] = [d for d in dirs if d not in ('node_modules', '.git', 'dist', 'build', 'venv', '__pycache__')]
            for file in files:
                file_path = os.path.join(root, file)
                content = None
                
                if file.endswith('.md'):
                    try:
                        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                            content = f.read()
                    except Exception as e:
                        print(f"Error reading {file_path}: {e}")
                elif file.endswith('.docx') and not file.startswith('~$'):
                    try:
                        content = extract_docx_text(file_path)
                    except Exception as e:
                        print(f"Error reading {file_path}: {e}")
                        
                if content and content.strip():
                    parts = os.path.normpath(file_path).split(os.sep)
                    folder_category = parts[-2] if len(parts) >= 2 else "general"
                    
                    doc = Document(
                        page_content=content.strip(),
                        metadata={
                            "source": file_path,
                            "filename": file,
                            "folder": folder_category,
                            "company": "relanto" if folder_category == "seller" else folder_category,
                            "document_type": file.rsplit('.', 1)[0]
                        }
                    )
                    documents.append(doc)
        return documents

